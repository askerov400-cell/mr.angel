"""Bounded, offline parsers. Invoked in a disposable child process."""
import base64
import csv
import io
import json
import sys
import zipfile
from pathlib import Path
from html.parser import HTMLParser

MAX_TEXT = 300000
MAX_ROWS = 2000
MAX_COLS = 100
EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.xls', '.csv', '.txt', '.md',
              '.json', '.html', '.htm', '.pptx', '.odt', '.jpg', '.jpeg', '.png', '.webp'}


class LimitError(ValueError):
    pass


class Lines:
    def __init__(self):
        self.parts, self.size, self.rows = [], 0, 0

    def add(self, value):
        value = str(value).strip()
        if not value:
            return
        self.size += len(value) + (1 if self.parts else 0)
        if self.size > MAX_TEXT:
            raise LimitError('Слишком много текста: предел 300000 символов.')
        self.parts.append(value)

    def row(self, label, values):
        values = list(values)
        if len(values) > MAX_COLS:
            raise LimitError('Слишком много столбцов: предел 100.')
        self.rows += 1
        if self.rows > MAX_ROWS:
            raise LimitError('Слишком много строк: предел 2000.')
        self.add(label + ': ' + json.dumps(values, ensure_ascii=False, default=str))


def archive(path):
    z = zipfile.ZipFile(path)
    if len(z.infolist()) > 2000 or sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
        z.close()
        raise LimitError('Слишком большой распакованный документ.')
    return z


def xml_text(node):
    return ''.join(node.itertext()).strip()


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.hidden += 1
        if tag in ('p', 'div', 'br', 'tr', 'li', 'h1', 'h2', 'td', 'th'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style') and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def decode(data):
    encodings = ('utf-16',) if data.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'cp1251')
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            if '\x00' not in text:
                return text
        except UnicodeError:
            pass
    raise ValueError('Unsupported encoding')


def extract(path, ext):
    out, images = Lines(), []
    if ext not in EXTENSIONS:
        raise ValueError('Unsupported format')
    if ext in ('.txt', '.md', '.json', '.csv', '.html', '.htm'):
        text = decode(Path(path).read_bytes())
        if ext == '.csv':
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            for n, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                out.row('Строка ' + str(n), row)
        elif ext in ('.html', '.htm'):
            parser = HTMLText()
            parser.feed(text)
            out.add(''.join(parser.parts))
        else:
            out.add(text)
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        with archive(path):
            pass
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            out.add('Формулы приведены как текст; результаты не пересчитывались.')
            for sheet in book:
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLS:
                    raise LimitError('Предел таблицы: 2000 строк, 100 столбцов.')
                for n, row in enumerate(sheet.iter_rows(), 1):
                    if any(cell.value is not None for cell in row):
                        out.row(sheet.title + '!строка ' + str(n),
                                [f'{cell.coordinate}={cell.value}' for cell in row if cell.value is not None])
        finally:
            book.close()
    elif ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            out.add('XLS: сохранённые значения ячеек, формулы не пересчитывались.')
            for sheet in book.sheets():
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLS:
                    raise LimitError('Предел таблицы: 2000 строк, 100 столбцов.')
                for n in range(sheet.nrows):
                    values = [xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
                              if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                              for cell in sheet.row(n)]
                    out.row(sheet.name + '!строка ' + str(n+1), values)
        finally:
            book.release_resources()
    elif ext in ('.docx', '.odt', '.pptx'):
        from defusedxml.ElementTree import fromstring
        with archive(path) as z:
            if ext == '.docx':
                root = fromstring(z.read('word/document.xml'))
                ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                for child in root.find(ns+'body'):
                    if child.tag == ns+'p':
                        out.add(''.join(child.itertext()))
                    elif child.tag == ns+'tbl':
                        for n, row in enumerate(child.findall(ns+'tr'), 1):
                            out.row('Таблица, строка '+str(n), [xml_text(c) for c in row.findall(ns+'tc')])
            elif ext == '.odt':
                root = fromstring(z.read('content.xml'))
                ns = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
                table_ns = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
                for node in root.iter():
                    if node.tag in (ns+'p', ns+'h'):
                        out.add(xml_text(node))
                    elif node.tag == table_ns+'table-row':
                        out.row('Таблица ODT, строка', [xml_text(c) for c in node.findall(table_ns+'table-cell')])
            else:
                names = [n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml') and '/_rels/' not in n]
                names.sort(key=lambda n: int(Path(n).stem.removeprefix('slide')))
                if len(names) > 40:
                    raise LimitError('Предел: 40 слайдов.')
                ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
                for n, name in enumerate(names, 1):
                    out.add('Слайд '+str(n))
                    root = fromstring(z.read(name))
                    for node in root.iter(ns+'p'):
                        out.add(''.join(t.text or '' for t in node.iter(ns+'t')))
                    for table in root.iter(ns+'tbl'):
                        for row in table.findall(ns+'tr'):
                            out.row('Слайд '+str(n)+', строка таблицы', [xml_text(c) for c in row.findall(ns+'tc')])
        out.add('[Прочитан текст; встроенные изображения не распознавались.]')
    elif ext in ('.jpg', '.jpeg', '.png', '.webp'):
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = 20000000
        with Image.open(path) as image:
            if image.width * image.height > 20000000:
                raise LimitError('Слишком большое разрешение изображения.')
            image = ImageOps.exif_transpose(image).convert('RGB')
            image.thumbnail((1800, 1800))
            stream = io.BytesIO()
            image.save(stream, format='JPEG', quality=90)
            images.append({'page':1, 'jpeg':base64.b64encode(stream.getvalue()).decode('ascii')})
            out.add('[Нужен OCR страницы 1]')
    else:
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise LimitError('Документ защищён паролем. Пришли незашифрованную копию.')
            if len(doc) > 40:
                raise LimitError('Предел PDF: 40 страниц.')
            for n, page in enumerate(doc, 1):
                text = page.get_text() if ext == '.pdf' else ''
                out.add('Страница '+str(n))
                if text.strip():
                    out.add(text)
                else:
                    if len(images) >= 5:
                        raise LimitError('Предел: 5 сканированных страниц. Раздели PDF.')
                    scale = min(2, 1800 / max(page.rect.width, page.rect.height))
                    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False)
                    images.append({'page': n, 'jpeg': base64.b64encode(pix.tobytes('jpeg')).decode('ascii')})
                    out.add('[Нужен OCR страницы '+str(n)+']')
    return {'text': '\n'.join(out.parts), 'images': images}


def main():
    try:
        if sys.platform != 'win32':
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (512*1024*1024, 512*1024*1024))
        result = extract(sys.argv[1], sys.argv[2])
    except LimitError as error:
        result = {'error': str(error)}
    except Exception:
        result = {'error': 'Не удалось прочитать файл: проверь формат и целостность.'}
    print(json.dumps(result, ensure_ascii=True))


if __name__ == '__main__':
    main()
